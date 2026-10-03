"""Minimal house deck kit (approximation of Glilot style until the G: drive deck_kit.py/template is available).
Style per DECK_BLUEPRINT.md: 16:9, pale #F4F8FE wash, navy #1F2D5A titles top-left with grey italic subhead,
periwinkle corner wedge, source footnote on every slide, titles <=70 chars, body >=14pt."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

C = lambda h: RGBColor.from_string(h.lstrip('#'))
NAVY, BLUE, BG, WEDGE, GREY, INK2 = C('1F2D5A'), C('1B6CB5'), C('F4F8FE'), C('C9D8F0'), C('8A94A6'), C('3F4A5C')
GOOD, AMBER, RED, WHITE = C('2E7D32'), C('E69500'), C('C0392B'), C('FFFFFF')
CAT = [C(x) for x in ('5B9BD5', 'F2B6A0', 'A8C97F', 'F2B23E', 'C9C2E0', '3FA39B', '1B6CB5')]
W, H = Inches(13.333), Inches(7.5)

class Deck:
    def __init__(self, footer):
        self.p = Presentation(); self.p.slide_width, self.p.slide_height = W, H; self.footer = footer; self.n = 0
        self.titles = []

    def slide(self, title, sub=None, source=None):
        assert len(title) <= 70, f'title too long ({len(title)}): {title}'
        s = self.p.slides.add_slide(self.p.slide_layouts[6]); self.n += 1; self.titles.append(title)
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = BG
        w = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, 0, H - Inches(1.1), Inches(1.1), Inches(1.1))
        w.fill.solid(); w.fill.fore_color.rgb = WEDGE; w.line.fill.background()
        self.text(s, title, Inches(.6), Inches(.35), Inches(12.1), Inches(.75), 28, NAVY, bold=True)
        if sub: self.text(s, sub, Inches(.6), Inches(1.05), Inches(12.1), Inches(.45), 15, GREY, italic=True)
        foot = (f'Source: {source}  ·  ' if source else '') + self.footer
        self.text(s, foot, Inches(1.2), H - Inches(.42), Inches(10.6), Inches(.3), 9, GREY, italic=True)
        self.text(s, str(self.n), W - Inches(.8), H - Inches(.42), Inches(.4), Inches(.3), 9, GREY, align=PP_ALIGN.RIGHT)
        return s

    def text(self, s, t, x, y, w, h, size=14, color=INK2, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(.05)
        lines = t if isinstance(t, list) else [t]
        for i, ln in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.alignment = align
            r = para.add_run(); r.text = ln; f = r.font; f.size = Pt(size); f.color.rgb = color; f.bold = bold; f.italic = italic; f.name = 'Calibri'
            para.space_after = Pt(10 if size >= 20 else 4)
        return tb

    def bullets(self, s, items, x, y, w, h, size=15, color=INK2):
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        for i, it in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            bold = None
            if isinstance(it, tuple): bold, it = it
            if bold:
                r = para.add_run(); r.text = bold + ' '; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY; r.font.name = 'Calibri'
            r = para.add_run(); r.text = ('' if bold else '• ') + it; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = 'Calibri'
            para.space_after = Pt(8)
        return tb

    def box(self, s, x, y, w, h, fill=WHITE, line=C('D7DDE6')):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h); b.adjustments[0] = .06
        b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = line; b.shadow.inherit = False
        return b

    def tile(self, s, x, y, w, h, label, value, note=None, color=NAVY):
        self.box(s, x, y, w, h)
        self.text(s, label.upper(), x + Inches(.15), y + Inches(.12), w - Inches(.3), Inches(.3), 11, GREY, bold=True)
        self.text(s, value, x + Inches(.15), y + Inches(.4), w - Inches(.3), Inches(.7), 28, color, bold=True)
        if note: self.text(s, note, x + Inches(.15), y + h - Inches(.55), w - Inches(.3), Inches(.5), 12, INK2)

    def card(self, s, x, y, w, h, head, body, color=BLUE, size=14):
        self.box(s, x, y, w, h)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(.08)); bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background()
        self.text(s, head, x + Inches(.15), y + Inches(.18), w - Inches(.3), Inches(.45), 16, color, bold=True)
        if isinstance(body, list): self.bullets(s, body, x + Inches(.12), y + Inches(.65), w - Inches(.25), h - Inches(.75), size)
        else: self.text(s, body, x + Inches(.15), y + Inches(.65), w - Inches(.3), h - Inches(.75), size)

    def image(self, s, path, x, y, w=None, h=None):
        return s.shapes.add_picture(path, x, y, width=w, height=h)

    def table(self, s, rows, x, y, w, h, size=13, col_w=None, header_fill=NAVY, heat=None):
        t = s.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
        if col_w:
            for i, cw in enumerate(col_w): t.columns[i].width = Inches(cw)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                cell = t.cell(i, j); cell.text = str(v); p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(size); p.font.name = 'Calibri'
                cell.margin_left = cell.margin_right = Inches(.06); cell.margin_top = cell.margin_bottom = Inches(.03)
                if i == 0:
                    cell.fill.solid(); cell.fill.fore_color.rgb = header_fill; p.font.color.rgb = WHITE; p.font.bold = True
                else:
                    cell.fill.solid(); cell.fill.fore_color.rgb = WHITE if i % 2 else C('EEF3FA'); p.font.color.rgb = INK2
                    if heat and (i, j) in heat: cell.fill.fore_color.rgb = heat[(i, j)]
                if j > 0 and len(str(v)) <= 16: p.alignment = PP_ALIGN.CENTER
        return t

    def bar_chart(self, s, cats, series, x, y, w, h, horizontal=False, fmt='0', legend=True, colors=None):
        cd = CategoryChartData(); cd.categories = cats
        for name, vals in series: cd.add_series(name, vals)
        ct = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        ch = s.shapes.add_chart(ct, x, y, w, h, cd).chart
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
        ch.value_axis.tick_labels.font.size = Pt(12); ch.category_axis.tick_labels.font.size = Pt(12)
        ch.value_axis.tick_labels.number_format = fmt; ch.value_axis.tick_labels.number_format_is_linked = False
        ch.value_axis.major_gridlines.format.line.color.rgb = C('E6EAF0')
        for i, ser in enumerate(ch.series):
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = (colors or CAT)[i % 7] if not (colors and len(series) == 1) else (colors[0])
            ser.data_labels.show_value = True; ser.data_labels.font.size = Pt(11); ser.data_labels.number_format = fmt; ser.data_labels.number_format_is_linked = False
        if colors and len(series) == 1 and len(colors) == len(cats):
            for i, pt in enumerate(ch.series[0].points): pt.format.fill.solid(); pt.format.fill.fore_color.rgb = colors[i]
        return ch

    def save(self, path):
        self.p.save(path); return path
