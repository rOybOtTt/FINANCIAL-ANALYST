"""Peer model (lighter kit): one tab per peer with 4 years of history, EV / mean EBITDA today and value per share
at the house band (Murata's own band used as the sector proxy until each peer has 10 years of history), plus a Comps sheet."""
import json, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
F = pd.read_csv('data/clean/financials.csv'); F = F[~F.fiscal_year.astype(str).str.startswith('Q')]
Cn = pd.read_csv('data/clean/consensus.csv'); MB = json.load(open('data/murata_band.json'))['band']['all']
NAMES = {'6981.T': 'Murata', '009150.KS': 'Samsung Electro-Mechanics', '6976.T': 'Taiyo Yuden', '6762.T': 'TDK', '2327.TW': 'Yageo', '6971.T': 'Kyocera', '2492.TW': 'Walsin Technology'}
NAVY = '1F2D5A'; hdr = Font(bold=True, color='FFFFFF'); hf = PatternFill('solid', fgColor=NAVY); inp = PatternFill('solid', fgColor='FFF4CC')
wb = Workbook(); cs = wb.active; cs.title = 'Comps'
cs['A1'] = 'MLCC peers: comps on the house method (EV / cycle-mean EBITDA)'; cs['A1'].font = Font(bold=True, size=15, color=NAVY)
cs['A2'] = 'Local currency, mn. Peer history = 4 fiscal years (Yahoo fundamentals); band = Murata own p25/median/p75 as sector proxy. Yellow = input.'
cols = ['Ticker', 'Name', 'Price', 'Shares (mn)', 'Net cash', 'EV', 'Mean EBITDA', 'EV / mean EBITDA', 'Value @ p25', 'Value @ median', 'Value @ p75', 'Price vs median value', 'Fwd P/E', 'Consensus TP']
for j, c in enumerate(cols, 1):
    x = cs.cell(4, j, c); x.font = hdr; x.fill = hf; x.alignment = Alignment(wrap_text=True, horizontal='center')
cs['P4'] = 'Band'; cs['P5'] = 'p25'; cs['Q5'] = MB['p25']; cs['P6'] = 'median'; cs['Q6'] = MB['p50']; cs['P7'] = 'p75'; cs['Q7'] = MB['p75']
for r in (5, 6, 7): cs[f'Q{r}'].fill = inp
for i, (t, n) in enumerate(NAMES.items()):
    c = Cn[Cn.ticker == t].set_index('metric')['value']; f = F[F.ticker == t].pivot_table(index='fiscal_year', columns='metric', values='value', aggfunc='first')
    ws = wb.create_sheet(t.replace('.', '_'))
    ws['A1'] = f'{n} ({t})'; ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = 'History from Yahoo fundamentals (see data/clean/financials.csv); price/consensus stockanalysis.com 2-Oct-2026.'
    fys = [y for y in f.index if not pd.isna(f.loc[y].get('ebitda', float('nan')))]
    ws.cell(4, 1, 'Metric').font = hdr; ws.cell(4, 1).fill = hf
    for j, y in enumerate(fys): x = ws.cell(4, 2 + j, y); x.font = hdr; x.fill = hf
    mets = ['revenue', 'operating_income', 'ebitda', 'net_income', 'free_cash_flow', 'capex', 'cash', 'total_debt']
    for k, m in enumerate(mets, 5):
        ws.cell(k, 1, m.replace('_', ' '))
        for j, y in enumerate(fys):
            v = f.loc[y].get(m); ws.cell(k, 2 + j, None if pd.isna(v) else float(v)).number_format = '#,##0'
    last = chr(ord('A') + len(fys)); first = 'B'
    ws.cell(13, 1, 'Operating margin')
    for j in range(len(fys)): col = chr(66 + j); ws[f'{col}13'] = f'={col}6/{col}5'; ws[f'{col}13'].number_format = '0.0%'
    price = float(c.get('price')); mcap = float(c.get('market_cap')); nc = float(c.get('net_cash', 0) or 0)
    rows = [('Price', price, True), ('Shares (mn) = mcap / price', round(mcap / price, 3), True), ('Net cash (vendor)', nc, True),
            ('EV', '=B16*B17-B18', False), ('Mean EBITDA', f'=AVERAGE({first}7:{last}7)', False), ('EV / mean EBITDA', '=B19/B20', False),
            ('Value/share @ band p25', '=(Comps!$Q$5*B20+B18)/B17', False), ('Value/share @ median', '=(Comps!$Q$6*B20+B18)/B17', False),
            ('Value/share @ p75', '=(Comps!$Q$7*B20+B18)/B17', False), ('Multiple implied by price', '=B21', False)]
    for k, (a, v, isin) in enumerate(rows, 16):
        ws.cell(k, 1, a).font = Font(bold=True); x = ws.cell(k, 2, v); x.number_format = '0.0' if 'EBITDA' in a and '/' in a else '#,##0.0' if k == 17 else '#,##0'
        if isin: x.fill = inp
    ws.column_dimensions['A'].width = 30
    for col in 'BCDEF': ws.column_dimensions[col].width = 14
    sh = t.replace('.', '_'); r = 5 + i
    vals = [t, n, f"='{sh}'!B16", f"='{sh}'!B17", f"='{sh}'!B18", f"='{sh}'!B19", f"='{sh}'!B20", f"='{sh}'!B21", f"='{sh}'!B22", f"='{sh}'!B23", f"='{sh}'!B24", f'=C{r}/J{r}-1', c.get('fwd_pe'), c.get('target_price')]
    for j, v in enumerate(vals, 1):
        x = cs.cell(r, j, None if (isinstance(v, float) and pd.isna(v)) else v)
        x.number_format = '0.0' if j in (8, 13) else '0%' if j == 12 else '#,##0'
cs.conditional_formatting.add('H5:H11', ColorScaleRule(start_type='min', start_color='E0F2E7', mid_type='percentile', mid_value=50, mid_color='FFFFFF', end_type='max', end_color='F8E3E1'))
cs.column_dimensions['B'].width = 26
for col in 'CDEFGHIJKLMN': cs.column_dimensions[col].width = 13
cs['A13'] = 'Note: Murata own-band value differs from its tab here because the house model uses 5 FYs from fact books and forecasts to FY3/30.'
wb.save('peers/MLCC_peers_model_v1_2.10.26.xlsx'); print('saved')
