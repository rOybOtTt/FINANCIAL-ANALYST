"""Recalculate the workbook in LibreOffice and print the key outputs; fails on any formula error."""
import subprocess, shutil, sys, os
from openpyxl import load_workbook
src = sys.argv[1]; tmp = '/tmp/lo/v.xlsx'; os.makedirs('/tmp/lo/out', exist_ok=True); shutil.copy(src, tmp)
subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'xlsx', '--outdir', '/tmp/lo/out', tmp], capture_output=True, timeout=180)
wb = load_workbook('/tmp/lo/out/v.xlsx', data_only=True)
err = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and (c.value.startswith('#') or c.value.startswith('Err:'))]
none = [(ws.title, c.coordinate) for ws in [wb['Valuation']] for row in ws.iter_rows(min_row=5, max_row=13, min_col=2, max_col=4) for c in row if c.value is None]
V = wb['Valuation']
for r in V.iter_rows(min_row=4, max_row=24, values_only=True): print([round(x, 3) if isinstance(x, float) else x for x in r[:5]])
G = wb['Sensitivity']
for r in G.iter_rows(min_row=5, max_row=11, values_only=True): print([round(x, 2) if isinstance(x, float) else x for x in r[:8]])
for r in G.iter_rows(min_row=16, max_row=24, values_only=True): print([round(x) if isinstance(x, float) else x for x in r[:8]])
print('band', [round(wb['Band'].cell(i, 6).value, 2) for i in range(5, 10)])
print('errors', err[:10], 'empty', none[:5]); print('PASS' if not err and not none else 'FAIL')
